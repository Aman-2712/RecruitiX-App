import os
import re
import json
import logging
from datetime import datetime
from typing import Dict, Any, Tuple, List
from pypdf import PdfReader
from docx import Document
from openai import OpenAI

logger = logging.getLogger("hirecue.ai_service")

# Setup NVIDIA NIM API Client & OpenAI Client
nvidia_api_key = os.getenv("NVIDIA_API_KEY", "")
NVIDIA_MODEL = os.getenv("NVIDIA_MODEL", "meta/llama-3.3-70b-instruct")

nvidia_client = None
if nvidia_api_key:
    nvidia_client = OpenAI(api_key=nvidia_api_key, base_url="https://integrate.api.nvidia.com/v1")

# Setup OpenAI Client
openai_api_key = os.getenv("OPENAI_API_KEY", "")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
client = None
if openai_api_key:
    client = OpenAI(api_key=openai_api_key)

def clean_extracted_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def extract_text_from_pdf(file_path: str) -> str:
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text += content + "\n"
        return clean_extracted_text(text)
    except Exception as e:
        logger.error(f"Error reading PDF {file_path}: {e}")
        return ""

def extract_text_from_docx(file_path: str) -> str:
    # Handles .docx files using python-docx
    try:
        doc = Document(file_path)
        text = []
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    text.append(" | ".join(cells))
        return clean_extracted_text("\n".join(text))
    except Exception as e:
        logger.error(f"Error reading DOCX {file_path}: {e}")
        return ""

def extract_text_from_doc(file_path: str) -> str:
    # Handles legacy .doc files when textract is available in the runtime.
    try:
        import textract
        raw = textract.process(file_path)
        return clean_extracted_text(raw.decode("utf-8", errors="ignore"))
    except Exception as e:
        logger.error(f"Error reading DOC {file_path}: {e}")
        return ""

def extract_resume_text(file_path: str) -> str:
    _, ext = os.path.splitext(file_path.lower())
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        if ext == ".docx":
            return extract_text_from_docx(file_path)
        else:
            return extract_text_from_doc(file_path)
    else:
        # Fallback text read
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return clean_extracted_text(f.read())
        except Exception as e:
            logger.error(f"Error reading raw file {file_path}: {e}")
            return ""

# --- MOCK RESUME PARSER & MATCHER (Fallback when no OpenAI API Key) ---

COMMON_SKILLS = [
    "Python", "JavaScript", "TypeScript", "React", "Next.js", "Vue", "Angular", 
    "Node.js", "Express", "FastAPI", "Flask", "Django", "Java", "C++", "C#", "Go", "Rust",
    "PostgreSQL", "MySQL", "MongoDB", "Redis", "SQLite", "SQL", "HTML", "CSS", "Tailwind",
    "AWS", "Azure", "GCP", "Docker", "Kubernetes", "Git", "GitHub", "CI/CD", "DevOps",
    "Machine Learning", "Data Science", "AI", "NLP", "PyTorch", "TensorFlow", "Pandas",
    # Sales, Marketing & PM
    "Cold Calling", "Lead Generation", "CRM", "Negotiation", "Pitching", "Sales", "Marketing",
    "HubSpot", "Google Analytics", "SEO", "SEM", "Content Strategy", "Copywriting",
    "Product Strategy", "User Research", "Agile", "Jira", "Figma",
    # Healthcare & Pharma
    "Nursing", "CPR", "ACLS", "BLS", "Patient Care", "Clinical", "Acute Care", "ICU", "Pharmacology", "Phlebotomy",
    # Finance, Accounting & Banking
    "Financial Analysis", "Accounting", "Tally", "Financial Modeling", "Auditing", "Taxation", "GST", "CPA", "CFA", "Risk Management", "Banking",
    # Legal & HR
    "Corporate Law", "Contract Drafting", "Legal Research", "Compliance", "Talent Acquisition", "Employee Relations", "Payroll", "HRMS",
    # Construction, Engineering & Manufacturing
    "AutoCAD", "Revit", "Civil Engineering", "Project Management", "PMP", "Six Sigma", "Quality Control", "OSHA", "Supply Chain", "Logistics",
    # Retail, Hospitality & Education
    "Customer Service", "Inventory Management", "Point of Sale", "Teaching", "Curriculum Development", "Public Speaking"
]


def extract_dates_and_calculate_experience(text: str) -> int:
    """
    Extracts date ranges from resume text and calculates total years of experience.
    """
    # Regex to find patterns like Jan 2020 - Dec 2022, 06/2019 to Present, 2018 - 2021, etc.
    months_pat = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
    month = r"(?:" + months_pat + r"|\d{1,2})"
    sep = r"[\s/-]+"
    year = r"\b(?:19\d{2}|20\d{2})\b"
    
    date_pat = r"(?:" + month + sep + r")?" + year
    range_pat = r"(" + date_pat + r")\s*(?:-|to|–|—|until)\s*CustomDatePat".replace("CustomDatePat", "(" + date_pat + r"|present|current|now|active)")
    
    matches = re.findall(range_pat, text, re.IGNORECASE)
    
    if not matches:
        years = [int(y) for y in re.findall(r"\b(?:19\d{2}|20\d{2})\b", text)]
        if len(years) >= 2:
            return max(1, max(years) - min(years))
        return 2  # default fallback
        
    total_months = 0
    from datetime import datetime
    current_year = datetime.now().year
    current_month = datetime.now().month
    
    def parse_single_date(date_str: str) -> tuple[int, int]:
        date_str = date_str.lower().strip()
        if date_str in ["present", "current", "now", "active", ""]:
            return current_year, current_month
            
        yr_match = re.search(r"\b(?:19\d{2}|20\d{2})\b", date_str)
        if not yr_match:
            return current_year, current_month
        yr = int(yr_match.group(0)) # group(0) since it's non-capturing now
        
        m_word_match = re.search(months_pat, date_str)
        if m_word_match:
            m_word = m_word_match.group(0)
            month_mapping = {
                "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
                "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12
            }
            return yr, month_mapping.get(m_word[:3], 1)
            
        m_num_match = re.search(r"\b(\d{1,2})\b", date_str.replace(str(yr), ""))
        if m_num_match:
            m = int(m_num_match.group(1))
            if 1 <= m <= 12:
                return yr, m
                
        return yr, 1

    for start_str, end_str in matches:
        try:
            start_yr, start_m = parse_single_date(start_str)
            end_yr, end_m = parse_single_date(end_str)
            
            diff_months = (end_yr - start_yr) * 12 + (end_m - start_m)
            if diff_months > 0:
                total_months += diff_months
        except Exception:
            pass
            
    years_exp = total_months / 12.0
    return max(0, min(50, int(round(years_exp))))

def mock_parse_resume(text: str, filename: str) -> Dict[str, Any]:
    # Extract Email
    email_match = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    
    # Extract Phone
    phone_match = re.search(r"(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{2,5}\)?[-.\s]?)?\d{3,5}[-.\s]?\d{4}", text)
    phone = phone_match.group(0).strip() if phone_match else None
    
    # Extract Name (first non-empty line, or filename basis)
    name = None
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if lines:
        for line in lines[:8]:
            lower = line.lower()
            if any(word in lower for word in ["resume", "curriculum", "email", "phone", "linkedin", "github"]):
                continue
            if len(line.split()) >= 2 and len(line) < 45 and not any(char in line for char in ["@", ":", "/", "\\", "|"]):
                name = line
                break
    if not name:
        # Extract from filename
        clean_name = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ")
        clean_name = re.sub(r"\d+", "", clean_name).strip()
        if len(clean_name.split()) >= 2:
            name = clean_name.title()
    if not name:
        name = "Unknown Candidate"

    # Derive realistic email if not found in PDF raw text
    if email_match:
        email = email_match.group(0)
    else:
        clean_user = name.lower().replace(" ", ".").replace("..", ".")
        email = f"{clean_user}@talentpool.io"

    # Extract Skills
    found_skills = []
    for skill in COMMON_SKILLS:
        if re.search(r"\b" + re.escape(skill) + r"\b", text, re.IGNORECASE):
            found_skills.append(skill)

    # Better experiences parser:
    experiences = []
    range_regex = r"(?:(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)|\d{1,2})?[\s/-]*(?:19\d{2}|20\d{2})\s*(?:-|to|–|—)\s*(?:(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)|\d{1,2}|present|current|now)?[\s/-]*(?:19\d{2}|20\d{2})?"
    
    COMMON_JOB_TITLES = [
        "Software Engineer", "Developer", "Programmer", "Architect", "Manager",
        "Analyst", "Consultant", "Designer", "Lead", "Director", "VP", "President",
        "Officer", "Representative", "Associate", "Intern", "Specialist", "Administrator",
        "Engineer", "Scientist", "Coordinator", "Recruiter"
    ]
    
    for i, line in enumerate(lines):
        if re.search(range_regex, line, re.IGNORECASE):
            date_match = re.search(range_regex, line, re.IGNORECASE)
            date_str = date_match.group(0)
            
            parts = re.split(r"(?:-|to|–|—)", date_str, flags=re.IGNORECASE)
            start_date = parts[0].strip()
            end_date = parts[1].strip() if len(parts) > 1 else "Present"
            
            clean_line = line.replace(date_str, "").strip(" ,;()[]{}|")
            
            title = None
            company = None
            
            at_match = re.search(r"([A-Za-z0-9\s&]+)\s+(?:at|@)\s+([A-Za-z0-9\s&]+)", clean_line, re.IGNORECASE)
            if at_match:
                title = at_match.group(1).strip()
                company = at_match.group(2).strip()
            else:
                search_candidates = [clean_line]
                if i > 0:
                    search_candidates.append(lines[i-1])
                if i < len(lines) - 1:
                    search_candidates.append(lines[i+1])
                    
                for cand in search_candidates:
                    for ct in COMMON_JOB_TITLES:
                        if ct.lower() in cand.lower():
                            title = cand.strip(" ,;()[]{}|")
                            break
                    if title:
                        break
            
            if not title:
                if 5 < len(clean_line) < 50:
                    title = clean_line
                else:
                    title = "Professional Experience"
                    
            if not company:
                if i > 0 and len(lines[i-1]) < 40 and not any(ct.lower() in lines[i-1].lower() for ct in COMMON_JOB_TITLES):
                    company = lines[i-1]
                else:
                    company = "Company"
                    
            desc_lines = []
            for j in range(i+1, min(i+5, len(lines))):
                if re.search(range_regex, lines[j], re.IGNORECASE) or any(sec in lines[j].lower() for sec in ["education", "skills", "projects", "certifications"]):
                    break
                desc_lines.append(lines[j])
                
            experiences.append({
                "title": title.title(),
                "company": company.strip(" .,-()[]{}|"),
                "start_date": start_date,
                "end_date": end_date,
                "description": "\n".join(desc_lines)
            })
            
            if len(experiences) >= 5:
                break

    # Better education parser:
    educations = []
    edu_keywords = ["university", "college", "institute", "school", "academy", "polytechnic"]
    degree_keywords = ["bachelor", "master", "doctor", "phd", "b.s.", "m.s.", "b.a.", "m.a.", "btech", "mtech", "b.e.", "b.c.a.", "m.c.a.", "diploma", "associate"]
    
    for i, line in enumerate(lines):
        lower_line = line.lower()
        if any(keyword in lower_line for keyword in edu_keywords) or any(deg in lower_line for deg in degree_keywords):
            grad_year_match = re.search(r"\b(19\d{2}|20\d{2})\b", line)
            grad_year = grad_year_match.group(1) if grad_year_match else ""
            
            institution = None
            for word in line.split(","):
                if any(keyword in word.lower() for keyword in edu_keywords):
                    institution = word.strip()
                    break
            if not institution:
                institution = line.strip()
                
            degree = None
            major = None
            
            search_lines = [line]
            if i > 0:
                search_lines.append(lines[i-1])
            if i < len(lines) - 1:
                search_lines.append(lines[i+1])
                
            for sl in search_lines:
                for deg in degree_keywords:
                    if deg in sl.lower():
                        degree = deg.upper()
                        major_match = re.search(rf"\b{re.escape(deg)}\b(?:\s+(?:of|in|science|arts))?\s+([A-Za-z0-9\s&]+)", sl, re.IGNORECASE)
                        if major_match:
                            major = major_match.group(1).strip()
                        break
                if degree:
                    break
                    
            if not degree:
                degree = "Degree"
            if not major:
                major = "General Studies"
                
            educations.append({
                "institution": institution.title(),
                "degree": degree,
                "major": major.title(),
                "graduation_year": grad_year
            })
            
            if len(educations) >= 3:
                break

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "skills": found_skills if found_skills else ["JavaScript", "Python", "SQL"],
        "experiences": experiences,
        "educations": educations
    }

def mock_match_resume(candidate_json: Dict[str, Any], job_data: Dict[str, Any]) -> Dict[str, Any]:
    # Calculate Skill Match
    req_skills = job_data.get("skills_required", [])
    pref_skills = job_data.get("skills_preferred", [])
    cand_skills = [s.lower() for s in candidate_json.get("skills", [])]
    
    if not req_skills:
        skill_score = 80
    else:
        matched_req = [s for s in req_skills if s.lower() in cand_skills]
        req_score = (len(matched_req) / len(req_skills)) * 100
        
        # Penalize -15 points per missing required skill
        missing_req_count = len(req_skills) - len(matched_req)
        
        matched_pref = [s for s in pref_skills if s.lower() in cand_skills]
        pref_score = (len(matched_pref) / len(pref_skills)) * 100 if pref_skills else 100
        
        skill_score = int(req_score * 0.7 + pref_score * 0.3)
        if missing_req_count > 0:
            skill_score = max(0, skill_score - (missing_req_count * 15))

    # Calculate Experience Match using date parser
    min_exp = job_data.get("min_experience", 0)
    
    # Estimate candidate experience from work history
    cand_exp_years = 0
    experiences_text = " ".join([
        f"{exp.get('title', '')} {exp.get('company', '')} {exp.get('start_date', '')} to {exp.get('end_date', '')} {exp.get('description', '')}"
        for exp in candidate_json.get("experiences", [])
    ])
    if experiences_text.strip():
        cand_exp_years = extract_dates_and_calculate_experience(experiences_text)
    
    # If still 0, try default estimation
    if cand_exp_years == 0:
        for exp in candidate_json.get("experiences", []):
            cand_exp_years += 2
            
    if cand_exp_years >= min_exp:
        exp_score = min(100, 85 + int((cand_exp_years - min_exp) * 5))
    else:
        # Penalize if below target min experience
        exp_score = max(0, 80 - int((min_exp - cand_exp_years) * 15))

    # Calculate Relevance Score
    job_title = job_data.get("title", "").lower()
    cand_titles = [exp.get("title", "").lower() for exp in candidate_json.get("experiences", [])]
    
    relevance_score = 60
    if any(job_title in title or title in job_title for title in cand_titles if title):
        relevance_score += 25
    desc = job_data.get("description", "").lower()
    text_to_search = " ".join([
        exp.get("description", "") for exp in candidate_json.get("experiences", [])
    ]).lower() + " " + " ".join(cand_skills)
    
    # Extract domain-agnostic meaningful words from JD (works for Healthcare, Legal, Finance, Tech, etc.)
    stopwords = {"with", "that", "this", "from", "have", "will", "your", "must", "work", "team", "their", "about", "which", "would", "there", "other", "should", "ability", "experience", "candidate", "role", "position"}
    jd_words = set(re.findall(r'\b[a-z]{4,}\b', desc)) - stopwords
    matches = [word for word in jd_words if word in text_to_search]
    relevance_score += min(15, len(matches) * 2)
    relevance_score = min(100, relevance_score)

    match_score = int(skill_score * 0.4 + exp_score * 0.4 + relevance_score * 0.2)

    recs = []
    concerns = []
    
    if skill_score > 80:
        recs.append("Strong technical skill match for core job requirements.")
    if cand_exp_years >= min_exp:
        recs.append(f"Meets or exceeds minimum experience requirements ({cand_exp_years} years vs {min_exp} years required).")
    
    matched_skills = [s for s in (req_skills + pref_skills) if s.lower() in cand_skills]
    if matched_skills:
        recs.append(f"Demonstrates expertise in key technologies: {', '.join(matched_skills[:4])}.")
        
    if not recs:
        recs.append("Matches the baseline qualifications for the position.")
        
    missing_req = [s for s in req_skills if s.lower() not in cand_skills]
    if missing_req:
        concerns.append(f"Missing experience in required skills: {', '.join(missing_req[:3])}.")
    if cand_exp_years < min_exp:
        concerns.append(f"Has {cand_exp_years} years of experience which is below the target {min_exp} years.")
    if relevance_score < 70:
        concerns.append("Project portfolio shows limited direct overlap with our industry domain.")
        
    if not concerns:
        concerns.append("No major red flags or missing skill requirements identified.")

    return {
        "match_score": match_score,
        "skill_match_score": skill_score,
        "experience_match_score": exp_score,
        "relevance_score": relevance_score,
        "ai_summary": "\n".join([f"- {r}" for r in recs]),
        "ai_concerns": "\n".join([f"- {c}" for c in concerns])
    }



# --- OPENAI LLM RESUME PARSER & MATCHER ---

PARSE_PROMPT = """
You are an expert AI resume parsing agent for a modern ATS platform.
Your task is to parse the raw resume text below and extract candidate information into a clean, structured JSON object.

Strict Rules for Parsing:
1. **Name Extraction**: Find the candidate's actual full name. This is usually the first large text block at the top of the resume. Do not use company names, headers (e.g., "Resume", "Curriculum Vitae"), or reference names.
2. **Contact Info**:
   - Extract the email address cleanly. Do not include trailing characters.
   - Format the phone number cleanly (e.g. "+1 (555) 123-4567" or "555-123-4567").
3. **Tech Stack & Skills**:
   - Do a deep scan of the entire resume: the dedicated "Skills" section, "Projects" descriptions, and "Work Experience" bullets.
   - Extract both hard technical skills (languages, frameworks, databases, tools, cloud services, methodologies) and core domain skills.
   - Normalize variants (e.g., "ReactJS", "React.js" -> "React"; "NodeJS", "Node" -> "Node.js"; "Postgres" -> "PostgreSQL"; "AWS services" -> "AWS").
4. **Work Experiences**:
   - For each job, extract: title, company, start_date, end_date, and description.
   - Standardize dates to "YYYY-MM" or "Month YYYY" (e.g. "2019-06" or "June 2019"). If a job is ongoing, use "Present".
   - Keep descriptions bulleted, concise, and focused on key achievements, technologies, and responsibilities.
5. **Education**:
   - Extract: institution (university/college), degree (e.g. Bachelor, Master, Associate), major (field of study), and graduation_year.
6. **Integrity**: Do not invent, hallucinate, or assume any information not explicitly present in the text. Return null for missing scalar fields and [] for empty arrays.

Target JSON schema output format:
{
  "name": "Full Name",
  "email": "email@example.com",
  "phone": "+1 (555) 123-4567",
  "skills": ["Python", "React", "PostgreSQL", "Docker", "Git"],
  "experiences": [
    {
      "title": "Software Engineer",
      "company": "Company Name",
      "start_date": "June 2019",
      "end_date": "Present",
      "description": "- Developed REST APIs using Python and FastAPI.\\n- Integrated PostgreSQL database and optimized query latency by 30%.\\n- Deployed services to AWS using Docker."
    }
  ],
  "educations": [
    {
      "institution": "University Name",
      "degree": "Bachelor of Science",
      "major": "Computer Science",
      "graduation_year": "2018"
    }
  ]
}

Resume Text:
\"\"\"
{text}
\"\"\"
"""

MATCH_PROMPT = """
You are a senior AI Technical Recruiter and Job Matching Agent. Compare the candidate's parsed profile against the job description requirements.
Determine how well the candidate fits the role.

Job Details:
- Title: {job_title}
- Location: {job_location}
- Min Experience: {job_min_exp} years
- Required Skills: {job_req_skills}
- Preferred Skills: {job_pref_skills}
- Description: {job_description}

Candidate Profile (JSON):
{candidate_json}

Strict Rules for Scoring & Analysis:
1. **Skill Match Score (0-100)**:
   - Identify which Required Skills and Preferred Skills the candidate actually possesses.
   - Required skills are critical; missing even one required skill must significantly penalize this score (e.g., -15 points per missing required skill).
   - Do not give credit for similar-sounding skills unless they are direct aliases (e.g., do not match "Java" if they only list "JavaScript").
2. **Experience Match Score (0-100)**:
   - Calculate the candidate's total years of experience by summing the duration of their work history from their job start and end dates.
   - If they have less than {job_min_exp} years, penalize the score proportionally (e.g. if {job_min_exp} is 5 and they have 3 years, score should be around 60/100).
   - If they exceed or meet the minimum experience, the score should be 85-100 depending on the quality and progression of their titles.
3. **Relevance Score (0-100)**:
   - Assess how relevant their past titles, responsibilities, and industry context are to this specific job.
   - A candidate whose past jobs are in a different domain (e.g., applying for a "Senior Backend Developer" but has only "Product Manager" or "Frontend Intern" experience) must receive a low relevance score (below 50).
4. **Overall Match Score (0-100)**:
   - Calculate the overall score strictly as the weighted average: (Skill Match * 0.4) + (Experience Match * 0.4) + (Relevance * 0.2). Round to the nearest integer.
5. **HR Copilot Recommendations & Gaps**:
   - **Why Recommended**: Provide 2-4 bullet points highlighting specific technical strengths, matching required skills, years of experience, and projects that directly prove their fit.
   - **Potential Concerns**: Provide 1-3 bullet points calling out concrete gaps: missing required/preferred skills, falling short of experience years, career gaps, or lack of domain-specific exposure. Be honest and critical; do not say "No concerns" if they are missing any required skills or experience.

Your output must be a valid JSON object matching this schema:
{
  "match_score": 88,
  "skill_match_score": 90,
  "experience_match_score": 85,
  "relevance_score": 90,
  "ai_summary": "- 6+ years of total software engineering experience, matching the senior requirement.\\n- Solid skills in React, TypeScript, and Next.js as required.\\n- Experience building production-ready apps and deploying them.",
  "ai_concerns": "- Lacks direct experience with AWS or cloud deployments (preferred skill).\\n- No backend experience listed in Python or PostgreSQL."
}
"""

def parse_resume_with_openai(text: str, ai_model: str = "GEMINI") -> Dict[str, Any]:
    if not client:
        return {}
    
    # Map branded agent names to actual valid OpenAI model identifiers
    # NEX (GEMINI) -> gpt-4o-mini (fast, cost-effective)
    # Aura-Sonnet (CLAUDE) -> gpt-4o (most capable)
    # Vortex (GPT) -> gpt-4o (full power)
    if ai_model == "CLAUDE":
        model_name = "gpt-4o"          # Aura-Sonnet 5.0 - premium accuracy
    elif ai_model == "GPT":
        model_name = "gpt-4o"          # Vortex-4o - full power GPT-4o
    else:  # GEMINI / NEX (default)
        model_name = "gpt-4o-mini"     # NEX - fast and efficient

    try:
        response = client.chat.completions.create(
            model=model_name,
            messages=[
                {"role": "system", "content": "You are a helpful ATS resume extractor. Return ONLY valid JSON."},
                {"role": "user", "content": PARSE_PROMPT.format(text=text)}
            ],
            response_format={"type": "json_object"},
            temperature=0.1
        )
        result_content = response.choices[0].message.content
        return json.loads(result_content)
    except Exception as e:
        logger.warning(f"AI parsing with model {model_name} failed: {e}. Falling back to gpt-4o-mini.")
        try:
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": "You are a helpful ATS resume extractor. Return ONLY valid JSON."},
                    {"role": "user", "content": PARSE_PROMPT.format(text=text)}
                ],
                response_format={"type": "json_object"},
                temperature=0.1
            )
            result_content = response.choices[0].message.content
            return json.loads(result_content)
        except Exception as fallback_err:
            logger.error(f"Fallback OpenAI parsing failed: {fallback_err}")
            return {}

def match_resume_with_openai(candidate_json: Dict[str, Any], job_data: Dict[str, Any], ai_model: str = "GEMINI") -> Dict[str, Any]:
    if not client:
        return {}
        
    # Same model mapping as above for consistency
    if ai_model == "CLAUDE":
        model_name = "gpt-4o"
    elif ai_model == "GPT":
        model_name = "gpt-4o"
    else:  # GEMINI / NEX
        model_name = "gpt-4o-mini"
        
    try:
        req_skills_str = ", ".join(job_data.get("skills_required", []))
        pref_skills_str = ", ".join(job_data.get("skills_preferred", []))
        prompt = MATCH_PROMPT.format(
            job_title=job_data.get("title", ""),
            job_location=job_data.get("location", ""),
            job_min_exp=job_data.get("min_experience", 0),
            job_req_skills=req_skills_str,
            job_pref_skills=pref_skills_str,
            job_description=job_data.get("description", ""),
            candidate_json=json.dumps(candidate_json)
        )
        response = client.chat.completions.create(
            model=model_name,
            messages=[
                {"role": "system", "content": "You are a helpful hiring manager matching agent. Return ONLY valid JSON."},
                {"role": "user", "content": prompt}
            ],
            response_format={"type": "json_object"},
            temperature=0.1
        )
        result_content = response.choices[0].message.content
        return json.loads(result_content)
    except Exception as e:
        logger.warning(f"AI matching with model {model_name} failed: {e}. Falling back to gpt-4o-mini.")
        try:
            req_skills_str = ", ".join(job_data.get("skills_required", []))
            pref_skills_str = ", ".join(job_data.get("skills_preferred", []))
            prompt = MATCH_PROMPT.format(
                job_title=job_data.get("title", ""),
                job_location=job_data.get("location", ""),
                job_min_exp=job_data.get("min_experience", 0),
                job_req_skills=req_skills_str,
                job_pref_skills=pref_skills_str,
                job_description=job_data.get("description", ""),
                candidate_json=json.dumps(candidate_json)
            )
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": "You are a helpful hiring manager matching agent. Return ONLY valid JSON."},
                    {"role": "user", "content": prompt}
                ],
                response_format={"type": "json_object"},
                temperature=0.1
            )
            result_content = response.choices[0].message.content
            return json.loads(result_content)
        except Exception as fallback_err:
            logger.error(f"Fallback OpenAI matching failed: {fallback_err}")
            return {}

def ensure_list(value: Any) -> List[Any]:
    if isinstance(value, list):
        return value
    if value is None:
        return []
    return [value]

def first_non_empty(*values: Any) -> Any:
    for value in values:
        if value not in (None, "", []):
            return value
    return None

def normalize_candidate_profile(profile: Dict[str, Any], text: str, filename: str) -> Dict[str, Any]:
    fallback = mock_parse_resume(text, filename)
    normalized = {
        "name": first_non_empty(profile.get("name"), fallback.get("name"), "Unknown Candidate"),
        "email": first_non_empty(profile.get("email"), fallback.get("email")),
        "phone": first_non_empty(profile.get("phone"), fallback.get("phone")),
        "skills": ensure_list(first_non_empty(profile.get("skills"), fallback.get("skills"))),
        "experiences": ensure_list(first_non_empty(profile.get("experiences"), fallback.get("experiences"))),
        "educations": ensure_list(first_non_empty(profile.get("educations"), fallback.get("educations"))),
    }
    normalized["skills"] = [
        str(skill).strip()
        for skill in normalized["skills"]
        if str(skill).strip()
    ]
    return normalized

def clamp_score(value: Any) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except Exception:
        return 0

def normalize_match_result(match_result: Dict[str, Any], fallback: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "match_score": clamp_score(first_non_empty(match_result.get("match_score"), fallback.get("match_score"))),
        "skill_match_score": clamp_score(first_non_empty(match_result.get("skill_match_score"), fallback.get("skill_match_score"))),
        "experience_match_score": clamp_score(first_non_empty(match_result.get("experience_match_score"), fallback.get("experience_match_score"))),
        "relevance_score": clamp_score(first_non_empty(match_result.get("relevance_score"), fallback.get("relevance_score"))),
        "ai_summary": str(first_non_empty(match_result.get("ai_summary"), fallback.get("ai_summary"), "")),
        "ai_concerns": str(first_non_empty(match_result.get("ai_concerns"), fallback.get("ai_concerns"), "")),
    }


# --- UNIFIED INTERFACE ---

def process_resume_and_match(file_path: str, filename: str, job_data: Dict[str, Any]) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """
    Extracts text from a resume, parses its details, and runs the matching comparison.
    Returns: (parsed_candidate_data, matched_scores_data)
    """
    ai_model = job_data.get("ai_model", "GEMINI")

    # 1. Extract Text
    raw_text = extract_resume_text(file_path)
    if not raw_text:
        # Fallback if parsing fails entirely
        raw_text = f"Empty resume text. Filename: {filename}"
        
    # 2. Parse Profile
    candidate_profile = {}
    if client:
        candidate_profile = parse_resume_with_openai(raw_text, ai_model)
        
    candidate_profile = normalize_candidate_profile(candidate_profile, raw_text, filename)
        
    # 3. Match against Job
    match_result = {}
    if client:
        match_result = match_resume_with_openai(candidate_profile, job_data, ai_model)
        
    fallback_match = mock_match_resume(candidate_profile, job_data)
    match_result = normalize_match_result(match_result, fallback_match)
        
    # Merge raw text back in
    candidate_profile["raw_text"] = raw_text
    
    return candidate_profile, match_result
